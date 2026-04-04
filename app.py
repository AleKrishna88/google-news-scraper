import requests
from bs4 import BeautifulSoup
from openai import OpenAI
from docx import Document
from io import BytesIO


# ======================
# GOOGLE NEWS RESULTS
# ======================

def get_google_news_results(keyword: str, num_results: int, serpapi_key: str, hl: str, gl: str):

    url = "https://serpapi.com/search.json"

    params = {
        "engine": "google_news",
        "q": keyword,
        "hl": hl,
        "gl": gl,
        "api_key": serpapi_key
    }

    response = requests.get(url, params=params, timeout=30)
    response.raise_for_status()

    data = response.json()

    news_results = data.get("news_results", [])

    competitors = []
    seen = set()

    blocked_domains = [
        "youtube.com",
        "youtu.be",
        "tiktok.com",
        "instagram.com",
        "facebook.com",
        "pinterest.com"
    ]

    for item in news_results:

        link = item.get("link")

        if not link:
            continue

        normalized_link = link.strip().rstrip("/")

        if normalized_link in seen:
            continue

        if any(domain in normalized_link for domain in blocked_domains):
            continue

        seen.add(normalized_link)

        competitors.append({
            "title": item.get("title", ""),
            "link": normalized_link,
            "source": item.get("source", "")
        })

        if len(competitors) >= num_results:
            break

    return competitors


# ======================
# PAGE SCRAPING
# ======================

def fetch_page(url: str):

    try:

        resp = requests.get(
            url,
            timeout=20,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/124.0.0.0 Safari/537.36"
                )
            }
        )

        resp.raise_for_status()

        html = resp.text

        soup = BeautifulSoup(html, "html.parser")

        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()

        text = " ".join(soup.get_text().split())

        return html, text[:18000]

    except Exception:
        return "", ""


def extract_metadata(html: str):

    soup = BeautifulSoup(html, "html.parser")

    title = soup.title.string.strip() if soup.title and soup.title.string else ""

    h1_tag = soup.find("h1")
    h1 = h1_tag.get_text(strip=True) if h1_tag else ""

    meta_desc = ""
    meta = soup.find("meta", attrs={"name": "description"})

    if meta and "content" in meta.attrs:
        meta_desc = meta["content"].strip()

    return title, h1, meta_desc


# ======================
# PARSE GPT OUTPUT
# ======================

def parse_generated_content(content: str):

    title = ""
    meta = ""
    article = content

    if "TITLE TAG:" in content and "META DESCRIPTION:" in content and "ARTICLE HTML:" in content:

        after_title = content.split("TITLE TAG:", 1)[1]
        title = after_title.split("META DESCRIPTION:", 1)[0].strip()

        after_meta = after_title.split("META DESCRIPTION:", 1)[1]
        meta = after_meta.split("ARTICLE HTML:", 1)[0].strip()

        article = after_meta.split("ARTICLE HTML:", 1)[1].strip()

    return title, meta, article


# ======================
# ARTICLE GENERATION
# ======================

def generate_article(keyword, competitors, openai_key, language):

    client = OpenAI(api_key=openai_key)

    merged = ""

    for comp in competitors:

        merged += f"""
URL: {comp['link']}

TITLE: {comp['html_title']}
H1: {comp['h1']}
META: {comp['meta_desc']}

CONTENUTO:
{comp['text']}

-------------------------
"""

    prompt = f"""
Sei un content writer SEO esperto.

Scrivi un contenuto SEO completo per la keyword:

{keyword}

Language code della ricerca: {language}

Il risultato deve contenere:

TITLE TAG (max 60 caratteri)
META DESCRIPTION (max 155 caratteri)
ARTICOLO HTML (800-1500 parole)

L'articolo deve essere scritto in HTML pronto per CMS.

Regole HTML:

- usa <h2> e <h3>
- usa <p>
- usa <ul> <ol>
- usa <strong>
- usa <table> se utile
- NON includere <html> <body>

Requisiti editoriali:

- Usa heading formulati come query di ricerca
- Testo discorsivo e informativo
- Inserisci elenchi quando utile
- Evidenzia le entità chiave con <strong>
- Evita contenuto di riempimento

Al termine inserisci almeno 4 FAQ in formato Q&A.

COMPETITOR DATA:
{merged}

Formato output:

TITLE TAG:
...

META DESCRIPTION:
...

ARTICLE HTML:
...
"""

    response = client.chat.completions.create(
        model="gpt-4.1",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.7
    )

    content = response.choices[0].message.content or ""

    return parse_generated_content(content)


# ======================
# WORD EXPORT
# ======================

def create_word_file(title_tag, meta_description, article):

    doc = Document()

    doc.add_heading("Title Tag", level=2)
    doc.add_paragraph(title_tag)

    doc.add_heading("Meta Description", level=2)
    doc.add_paragraph(meta_description)

    doc.add_heading("HTML Article", level=2)
    doc.add_paragraph(article)

    doc.save("seo_article.docx")

    print("\nFile salvato: seo_article.docx")


# ======================
# MAIN
# ======================

def main():

    keyword = input("Keyword: ")
    serpapi_key = input("SerpAPI key: ")
    openai_key = input("OpenAI key: ")

    language = "it"
    country = "it"
    num_results = 5

    print("\nRecupero Google News results...\n")

    competitors = get_google_news_results(
        keyword,
        num_results,
        serpapi_key,
        language,
        country
    )

    if not competitors:
        print("Nessun risultato trovato.")
        return

    enriched = []

    for comp in competitors:

        print("Scraping:", comp["link"])

        html, text = fetch_page(comp["link"])

        html_title, h1, meta_desc = extract_metadata(html)

        enriched.append({
            **comp,
            "html_title": html_title,
            "h1": h1,
            "meta_desc": meta_desc,
            "text": text
        })

    print("\nGenerazione articolo...\n")

    title_tag, meta_description, article = generate_article(
        keyword,
        enriched,
        openai_key,
        language
    )

    print("\nTITLE TAG:\n", title_tag)
    print("\nMETA DESCRIPTION:\n", meta_description)

    create_word_file(title_tag, meta_description, article)


if __name__ == "__main__":
    main()
